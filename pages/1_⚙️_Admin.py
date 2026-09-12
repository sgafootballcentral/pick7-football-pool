import streamlit as st
import requests
import pandas as pd
import time as time_module
import io
import uuid
from datetime import datetime, time, timezone
from zoneinfo import ZoneInfo
from supabase import create_client, Client
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

# 1. Connection and Secrets Verification
SUPABASE_URL = st.secrets.get("SUPABASE_URL")
SUPABASE_KEY = st.secrets.get("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    st.error("Missing critical Supabase connection variables in Streamlit Secrets.")
    st.stop()

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

st.title("⚙️ League Admin Panel")


def nudge_off_whole_number(odds_string: str) -> str:
    """'PHI -6' -> 'PHI -5.5' -- shifts a whole-number spread half a point toward
    zero so a final margin can never land exactly on the spread (no pushes)."""
    if not odds_string or odds_string == "0.0" or " " not in odds_string:
        return odds_string
    team_abbr, number_str = odds_string.rsplit(" ", 1)
    try:
        value = float(number_str)
    except ValueError:
        return odds_string
    if value != 0 and value == int(value):
        value += 0.5 if value < 0 else -0.5
    return f"{team_abbr} {value:.1f}"


def compute_game_numbers(week_games):
    """Same odd/even numbering as the weekly slate export: favorite=odd,
    underdog=even, assigned in chronological kickoff order."""
    sorted_games = sorted(week_games, key=lambda g: g.get("kickoff_time") or "")
    numbers = {}
    for i, g in enumerate(sorted_games, start=1):
        numbers[g["game_id"]] = {"fav_num": 2 * i - 1, "und_num": 2 * i}
    return numbers


def colored_picks_richtext(picks_nums, win_nums, loss_nums):
    """A single cell's worth of comma-separated picks, with each number colored
    green (bold) if it covered, red with a strikethrough if it didn't -- the
    closest robust analog to a hand-drawn circle/slash that survives resizing,
    re-opening, and different spreadsheet apps. Black if not graded yet.
    IMPORTANT: never set .font on a cell holding this -- it silently wipes the
    rich text back to None. Style with InlineFont here instead."""
    win_set, loss_set = set(win_nums), set(loss_nums)
    green = InlineFont(color="FF008000", b=True)
    red = InlineFont(color="FFFF0000", b=True, strike=True)
    black = InlineFont(color="FF000000", b=True)
    blocks = []
    for i, n in enumerate(picks_nums):
        if i > 0:
            blocks.append(TextBlock(black, ","))
        font = green if n in win_set else red if n in loss_set else black
        blocks.append(TextBlock(font, str(n)))
    return CellRichText(blocks)


def section_week_selector(key_prefix, master_week, label="Week:"):
    """A per-section week picker that follows the master 'Target Grouping Week
    Number' above by default. Checking the override box is the deliberate
    confirmation step to pick a different week just for this one section."""
    override = st.checkbox(
        f"Use a different week for this section (currently following Week {int(master_week)})",
        key=f"{key_prefix}_override",
    )
    if override:
        options = list(range(1, 17))
        default_index = options.index(int(master_week)) if int(master_week) in options else 0
        return st.selectbox(label, options=options, index=default_index, key=f"{key_prefix}_week_value")
    st.caption(f"Following the master week selector above: **Week {int(master_week)}**")
    return int(master_week)


@st.dialog("⚠️ Confirm Removal")
def confirm_removal(info):
    st.write(f"Remove **{info['player']}**'s picks for **Week {info['week']}**? This cannot be undone.")

    col_yes, col_no = st.columns(2)
    with col_yes:
        remove_clicked = st.button("Yes, remove", type="primary", use_container_width=True)
    with col_no:
        cancel_clicked = st.button("Cancel", use_container_width=True)

    if remove_clicked:
        if not info.get("user_id"):
            st.error("Couldn't find that player's user ID -- nothing was removed.")
        else:
            try:
                supabase.table("picks").delete().eq("user_id", info["user_id"]).eq("week_number", info["week"]).execute()
                st.session_state.pending_removal = None
                st.rerun()
            except Exception as e:
                st.error(f"Database error: {e}")
    elif cancel_clicked:
        st.session_state.pending_removal = None
        st.rerun()


@st.dialog("⚠️ Delete Entire Week?")
def confirm_week_delete(info):
    detail = f"This will permanently delete **{info['games_count']} game(s)** for Week {info['week']}"
    if info["picks_count"]:
        detail += f", along with **{info['picks_count']} pick(s)** from **{info['players_count']} player(s)**"
    detail += ". This cannot be undone."
    st.write(detail)

    col_yes, col_no = st.columns(2)
    with col_yes:
        confirm_clicked = st.button(f"Yes, delete Week {info['week']}", type="primary", use_container_width=True)
    with col_no:
        cancel_clicked = st.button("Cancel", use_container_width=True)

    if confirm_clicked:
        try:
            supabase.table("picks").delete().eq("week_number", info["week"]).execute()
            supabase.table("games").delete().eq("week_number", info["week"]).execute()
            st.session_state.pending_week_delete = None
            st.rerun()
        except Exception as e:
            st.error(f"Database error: {e}")
    elif cancel_clicked:
        st.session_state.pending_week_delete = None
        st.rerun()


@st.dialog("⚠️ Confirm Merge")
def confirm_merge(info):
    if not info.get("completed"):
        st.write(
            f"This will move **{info['from_picks_count']} pick(s)** from **{info['from_name']}** "
            f"into **{info['to_name']}**, then remove {info['from_name']}'s separate entry."
        )
        if info["conflict_weeks"]:
            st.warning(
                f"Week(s) {info['conflict_weeks']} already have picks under {info['to_name']} for the same "
                "game(s) -- those weeks will be SKIPPED to avoid duplicates. You'll need to sort those out by hand."
            )
        st.write("This cannot be undone.")

        col_yes, col_no = st.columns(2)
        with col_yes:
            confirm_clicked = st.button("Yes, merge", type="primary", use_container_width=True)
        with col_no:
            cancel_clicked = st.button("Cancel", use_container_width=True)

        if confirm_clicked:
            try:
                from_picks = supabase.table("picks").select("*").eq("user_id", info["from_id"]).execute().data
                conflict_weeks = set(info["conflict_weeks"])
                moved = 0
                for p in from_picks:
                    if p["week_number"] in conflict_weeks:
                        continue
                    supabase.table("picks").update({
                        "user_id": info["to_id"], "username": info["to_name"],
                    }).eq("id", p["id"]).execute()
                    moved += 1

                remaining = supabase.table("picks").select("id").eq("user_id", info["from_id"]).execute().data
                removed_entry = False
                if not remaining:
                    supabase.table("players").delete().eq("id", info["from_id"]).execute()
                    removed_entry = True

                info["completed"] = True
                info["moved"] = moved
                info["removed_entry"] = removed_entry
                st.session_state.pending_merge = info
                st.rerun()
            except Exception as e:
                st.error(f"Database error: {e}")
        elif cancel_clicked:
            st.session_state.pending_merge = None
            st.rerun()
    else:
        st.success(f"Moved {info['moved']} pick(s) into {info['to_name']}.")
        if not info.get("removed_entry"):
            st.warning(f"{info['from_name']}'s entry was kept because some weeks were skipped due to conflicts.")
        if st.button("Close", use_container_width=True):
            st.session_state.pending_merge = None
            st.rerun()


@st.dialog("⚠️ Delete This Player?")
def confirm_delete_player(info):
    detail = f"This will permanently delete **{info['name']}**"
    if info["picks_count"]:
        detail += f", along with **{info['picks_count']} pick(s)** across **{info['weeks_count']} week(s)**"
    detail += ". This cannot be undone."
    st.write(detail)

    col_yes, col_no = st.columns(2)
    with col_yes:
        confirm_clicked = st.button("Yes, delete", type="primary", use_container_width=True)
    with col_no:
        cancel_clicked = st.button("Cancel", use_container_width=True)

    if confirm_clicked:
        try:
            supabase.table("picks").delete().eq("user_id", info["player_id"]).execute()
            supabase.table("players").delete().eq("id", info["player_id"]).execute()
            st.session_state.pending_player_delete = None
            st.rerun()
        except Exception as e:
            st.error(f"Database error: {e}")
    elif cancel_clicked:
        st.session_state.pending_player_delete = None
        st.rerun()


@st.dialog("✍️ Confirm Picks Submission")
def confirm_manual_picks_save(info):
    st.write(f"You are submitting picks for **{info['player']}** — **Week {info['week']}**.")
    st.dataframe(info["preview_rows"], use_container_width=True, hide_index=True)

    col_yes, col_no = st.columns(2)
    with col_yes:
        confirm_clicked = st.button("Yes, save", type="primary", use_container_width=True)
    with col_no:
        cancel_clicked = st.button("Cancel", use_container_width=True)

    if confirm_clicked:
        try:
            supabase.table("picks").delete().eq("user_id", info["user_id"]).eq("week_number", info["week"]).execute()
            for item in info["resolved"]:
                supabase.table("picks").insert({
                    "user_id": info["user_id"],
                    "username": info["player"],
                    "week_number": info["week"],
                    "game_id": item["game_id"],
                    "selected_team": item["team"],
                    "spread_at_pick": item["spread"],
                }).execute()
            st.session_state.pending_manual_picks_save = None
            st.rerun()
        except Exception as e:
            st.error(f"Database error: {e}")
    elif cancel_clicked:
        st.session_state.pending_manual_picks_save = None
        st.rerun()

# 2. USER ACCESSIBILITY ROLE DATABASE CHECK
if "user" not in st.session_state or not st.session_state.user:
    st.warning("Please log in on the home page first.")
    st.stop()

if st.session_state.get("access_token"):
    try:
        supabase.auth.set_session(st.session_state.access_token, st.session_state.refresh_token)
    except Exception:
        st.session_state.user = None
        st.session_state.access_token = None
        st.warning("Your session expired -- please log in again on the home page.")
        st.stop()

user_id = st.session_state.user.id
is_admin = False

try:
    response = supabase.table("league_users").select("role").eq("id", user_id).execute()
    current_user_records = response.data
    if current_user_records and len(current_user_records) > 0:
        if current_user_records[0].get("role") == "admin":
            is_admin = True
except Exception as e:
    st.error(f"Database security check failed: {e}")
    is_admin = False

if not is_admin:
    st.error("🚫 Access Denied. Your profile role must be set to 'admin' in your database.")
    st.stop()

st.success("🔓 Commissioner Dashboard Unlocked!")

# Bidirectional week sync with the main app page (and Leaderboard), via a single
# shared "global_week" value. IMPORTANT: Streamlit deletes a widget's own
# session_state key entirely when you navigate away from the page it's on, so
# the widget key ("active_week") must be reseeded from global_week whenever
# it's missing -- not just "when global_week last changed" (that was the actual
# bug: the key can vanish from navigation even when global_week hasn't changed
# at all, and the old code only checked the latter).
if "global_week" not in st.session_state:
    st.session_state["global_week"] = 1

if "active_week" not in st.session_state:
    st.session_state["active_week"] = st.session_state["global_week"]

active_week = st.selectbox(
    "Target Grouping Week Number (For Player Submissions):",
    options=list(range(1, 17)), key="active_week",
)

if active_week != st.session_state["global_week"]:
    st.session_state["global_week"] = active_week


tab_setup, tab_picks, tab_grading, tab_players, tab_exports = st.tabs([
    "🗓️ Weekly Setup", "📋 Picks", "🏁 Grading", "👥 Players", "📤 Exports",
])

with tab_setup:
    # 📅 3. CALENDAR DATE SELECTOR
    st.subheader("📆 Select Custom Game Extraction Windows")
    selected_range = st.date_input(
        "Select Date Boundaries:",
        value=(datetime.today().date(), datetime.today().date()),
        help="Choose two specific calendar dates to set your pool window."
    )

    st.write("---")

    # 4. CHRONOLOGICAL DATE-DRIVEN AUTOMATED SYNCER
    def run_espn_sync(target_week: int, start_date, end_date):
        start_utc_str = datetime.combine(start_date, time.min).replace(tzinfo=timezone.utc).strftime("%Y%m%d")
        end_utc_str = datetime.combine(end_date, time.max).replace(tzinfo=timezone.utc).strftime("%Y%m%d")

        with st.spinner(f"Downloading game schedules from {start_date} to {end_date}..."):
            try:
                # Wipe previous entries for the selected grouping week to avoid duplication
                supabase.table("games").delete().eq("week_number", target_week).execute()

                date_range = f"{start_utc_str}-{end_utc_str}"
                leagues_to_fetch = [
                    {
                        "name": "NFL",
                        "url": "https://site.web.api.espn.com/apis/site/v2/sports/football/nfl/scoreboard",
                        "params": {"limit": 1000, "dates": date_range},
                    },
                    {
                        "name": "CFB",
                        "url": "https://site.web.api.espn.com/apis/site/v2/sports/football/college-football/scoreboard",
                        "params": {"limit": 1000, "groups": 80, "dates": date_range},
                    },
                ]

                total_games_inserted = 0
                skipped_no_line = 0
                headers = {
                    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
                    "Referer": "https://www.espn.com/",
                    "Accept": "application/json",
                }

                for target_league in leagues_to_fetch:
                    api_call = requests.get(target_league["url"], params=target_league["params"], headers=headers, timeout=15)
                    if api_call.status_code != 200:
                        st.warning(f"{target_league['name']} request failed: HTTP {api_call.status_code}")
                        continue
                    time_module.sleep(2)  # ESPN's hidden API has been rate-limiting back-to-back requests

                    response_data = api_call.json()

                    for event in response_data.get("events", []):
                        game_id = event.get("id")
                        kickoff_time = event.get("date")

                        # Target index array zero nodes safely
                        competitions = event.get("competitions", [{}])[0]
                        competitors = competitions.get("competitors", [])

                        # FIX: Safely read list array nodes for opening Vegas spreads
                        odds_array = competitions.get("odds", [])
                        has_line = bool(odds_array) and isinstance(odds_array, list)

                        home_node = next((c for c in competitors if c.get("homeAway") == "home"), competitors[0])
                        away_node = next((c for c in competitors if c.get("homeAway") == "away"), competitors[1])

                        home_team = home_node.get("team", {}).get("displayName", "Home Team")
                        away_team = away_node.get("team", {}).get("displayName", "Away Team")
                        home_abbr = home_node.get("team", {}).get("abbreviation", "").strip().upper()

                        broadcasts = competitions.get("broadcasts") or event.get("broadcasts") or []
                        tv_names = broadcasts[0].get("names", []) if broadcasts and isinstance(broadcasts, list) else []
                        tv_network = ", ".join(tv_names)

                        if not has_line:
                            skipped_no_line += 1
                            continue  # no book has posted a spread for this one -- leave it out of the pool

                        odds_string = odds_array[0].get("details", "0.0")

                        fav_team, und_team = away_team, home_team
                        fav_home, und_home = False, True

                        if odds_string != "0.0" and " " in odds_string:
                            line_parts = odds_string.split(" ")
                            fav_abbr_extracted = line_parts[0].strip().upper()

                            if fav_abbr_extracted == home_abbr:
                                fav_team, und_team = home_team, away_team
                                fav_home, und_home = True, False

                        # A genuine POSTED pick'em (0-point) line -- keep the game, just default
                        # the home team to -0.5 so it can't push.
                        spread_is_zero = True
                        if odds_string != "0.0" and " " in odds_string:
                            try:
                                spread_is_zero = float(odds_string.rsplit(" ", 1)[1]) == 0
                            except ValueError:
                                spread_is_zero = True
                        if spread_is_zero:
                            fav_team, und_team = home_team, away_team
                            fav_home, und_home = True, False
                            odds_string = f"{home_abbr} -0.5"

                        total_games_inserted += 1

                        # Save clean rows straight to your Supabase tables
                        supabase.table("games").insert({
                            "game_id": f"espn_{game_id}",
                            "game_number": total_games_inserted,
                            "league": target_league["name"],
                            "favorite_team": fav_team,
                            "underdog_team": und_team,
                            "favorite_team_home": fav_home,
                            "underdog_team_home": und_home,
                            "spread_value": nudge_off_whole_number(odds_string),
                            "display_text": f"{away_team} at {home_team}",
                            "kickoff_time": kickoff_time,
                            "tv_network": tv_network,
                            "week_number": target_week
                        }).execute()

                st.success(
                    f"Success! Pulled {total_games_inserted} games with a posted line "
                    f"(skipped {skipped_no_line} with no spread) from {start_date} to {end_date} into Week {target_week}!"
                )
                st.session_state.pending_refetch = None
                st.rerun()
            except Exception as e:
                st.error(f"ESPN Date Sync Failed: {e}")


    if st.button("🔄 Auto-Fetch Games by Selected Dates", type="primary"):
        if not (isinstance(selected_range, tuple) and len(selected_range) == 2):
            st.error("Validation Error: Please select both a start date and an end date on the calendar.")
        else:
            start_date, end_date = selected_range
            existing_picks = supabase.table("picks").select("user_id").eq("week_number", int(active_week)).execute().data
            if existing_picks:
                st.session_state.pending_refetch = {
                    "week": int(active_week),
                    "start": start_date,
                    "end": end_date,
                    "player_count": len(set(p["user_id"] for p in existing_picks)),
                }
                st.rerun()
            else:
                run_espn_sync(int(active_week), start_date, end_date)

    pending = st.session_state.get("pending_refetch")
    if pending and pending["week"] == int(active_week):
        st.warning(
            f"⚠️ Spreads for Week {pending['week']} have already been locked in by {pending['player_count']} "
            "player(s)' picks. Re-pulling is safe for them -- each pick keeps grading against the exact spread "
            "it locked in, even if the number changes here. This only affects the spread shown to anyone who "
            "hasn't submitted yet."
        )
        col_go, col_cancel = st.columns(2)
        with col_go:
            if st.button("Yes, re-pull anyway", type="primary", key="confirm_refetch"):
                run_espn_sync(pending["week"], pending["start"], pending["end"])
        with col_cancel:
            if st.button("Cancel", key="cancel_refetch"):
                st.session_state.pending_refetch = None
                st.rerun()

    st.write("---")

    # 5. ADD A GAME MANUALLY
    st.subheader("➕ Add a Game Manually")
    st.caption("For games ESPN doesn't have, or a line you want to set yourself.")

    manual_week = section_week_selector("manual_week", active_week, "Week number:")

    with st.form("manual_add_game_form"):
        col1, col2 = st.columns(2)
        with col1:
            manual_fav_team = st.text_input("Favorite team name")
        with col2:
            manual_und_team = st.text_input("Underdog team name")

        manual_home_side = st.radio("Who's the home team?", ["Favorite", "Underdog"], horizontal=True)
        manual_spread = st.number_input("Spread (favorite's number, e.g. -6.5):", value=-3.0, step=0.5, format="%.1f")

        col3, col4 = st.columns(2)
        with col3:
            manual_kickoff_date = st.date_input("Kickoff date", key="manual_kickoff_date")
        with col4:
            manual_kickoff_time = st.time_input("Kickoff time (Eastern)", key="manual_kickoff_time")

        manual_tv = st.text_input("TV network (optional)")

        manual_submitted = st.form_submit_button("Add Game")

    if manual_submitted:
        if not manual_fav_team.strip() or not manual_und_team.strip():
            st.error("Please enter both team names.")
        else:
            fav_team, und_team = manual_fav_team.strip(), manual_und_team.strip()
            fav_home = manual_home_side == "Favorite"
            spread_num = manual_spread

            if spread_num == 0:
                # Match the auto-fetch convention: a true pick'em defaults to the
                # home team as a -0.5 favorite so it can never push.
                if not fav_home:
                    fav_team, und_team = und_team, fav_team
                    fav_home = True
                spread_str = f"{fav_team} -0.5"
            else:
                spread_str = nudge_off_whole_number(f"{fav_team} {spread_num:.1f}")

            kickoff_dt_eastern = datetime.combine(manual_kickoff_date, manual_kickoff_time, tzinfo=ZoneInfo("America/New_York"))
            kickoff_iso = kickoff_dt_eastern.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")

            try:
                supabase.table("games").insert({
                    "game_id": f"manual_{uuid.uuid4().hex[:12]}",
                    "game_number": 0,
                    "league": "MANUAL",
                    "favorite_team": fav_team,
                    "underdog_team": und_team,
                    "favorite_team_home": fav_home,
                    "underdog_team_home": not fav_home,
                    "spread_value": spread_str,
                    "display_text": f"{und_team if fav_home else fav_team} at {fav_team if fav_home else und_team}",
                    "kickoff_time": kickoff_iso,
                    "tv_network": manual_tv.strip(),
                    "week_number": int(manual_week),
                }).execute()
                st.success(f"Added {fav_team} vs {und_team} to Week {manual_week}.")
                st.rerun()
            except Exception as e:
                st.error(f"Database error: {e}")

    st.write("---")

    # 10. DELETE WEEK'S SLATE
    st.subheader("🗑️ Delete Week's Slate")
    delete_week_num = section_week_selector("delete_week", active_week, "Week to delete:")

    if st.button("Delete This Week's Games", key="delete_week_btn"):
        games_for_delete = supabase.table("games").select("id").eq("week_number", delete_week_num).execute().data
        picks_for_delete = supabase.table("picks").select("user_id").eq("week_number", delete_week_num).execute().data

        st.session_state.pending_week_delete = {
            "week": delete_week_num,
            "games_count": len(games_for_delete),
            "picks_count": len(picks_for_delete),
            "players_count": len(set(p["user_id"] for p in picks_for_delete)),
        }
        st.rerun()

    if st.session_state.get("pending_week_delete"):
        confirm_week_delete(st.session_state.pending_week_delete)

    st.write("---")


with tab_picks:
    # 6. VIEW SUBMITTED PICKS
    st.subheader("📋 Submitted Picks")
    view_week = section_week_selector("view_week_picks", active_week, "Week to view picks for:")

    picks_rows = supabase.table("picks").select("*").eq("week_number", view_week).execute().data
    games_rows = supabase.table("games").select("game_id, display_text, favorite_team, underdog_team, spread_value, kickoff_time") \
        .eq("week_number", view_week).execute().data
    game_lookup = {g["game_id"]: g for g in games_rows}
    pick_numbers_map = compute_game_numbers(games_rows)

    all_players_rows = supabase.table("players").select("username").execute().data
    total_players = len(all_players_rows)

    picks_per_username = {}
    for p in picks_rows:
        uname = p.get("username", "Unknown")
        picks_per_username[uname] = picks_per_username.get(uname, 0) + 1

    submitted_count = len(picks_per_username)
    complete_count = sum(1 for n in picks_per_username.values() if n == 7)
    not_started_count = total_players - submitted_count

    col_total, col_submitted, col_complete, col_not_started = st.columns(4)
    col_total.metric("Total Players", total_players)
    col_submitted.metric("Submitted (any)", submitted_count)
    col_complete.metric("Complete (7/7)", complete_count)
    col_not_started.metric("Not Started", not_started_count)

    def status_for(n):
        if n == 0:
            return "❌ Not Started"
        if n == 7:
            return "✅ Complete"
        return f"⏳ {n}/7"

    roster_status = pd.DataFrame([
        {
            "Player": p["username"],
            "Picks Submitted": picks_per_username.get(p["username"], 0),
            "Status": status_for(picks_per_username.get(p["username"], 0)),
        }
        for p in all_players_rows
    ]).sort_values("Player").reset_index(drop=True)
    st.dataframe(roster_status, use_container_width=True, hide_index=True)

    if not picks_rows:
        st.info(f"No picks submitted yet for Week {view_week}.")
    else:
        display_rows = []
        for p in picks_rows:
            g = game_lookup.get(p["game_id"], {})
            nums = pick_numbers_map.get(p["game_id"], {})
            is_fav = p.get("selected_team") == g.get("favorite_team")
            pick_num = nums.get("fav_num") if is_fav else nums.get("und_num")
            display_rows.append({
                "Player": p.get("username", "Unknown"),
                "#": pick_num,
                "Matchup": g.get("display_text", p["game_id"]),
                "Pick": p.get("selected_team"),
                "Spread": g.get("spread_value", ""),
            })
        df_picks_view = pd.DataFrame(display_rows)

        players = sorted(df_picks_view["Player"].unique())
        selected_player = st.selectbox("View picks for:", ["— Select a player —"] + players, key="selected_picks_player")

        if selected_player == "— Select a player —":
            st.info("Select a player above to see their picks.")
        else:
            player_df = df_picks_view[df_picks_view["Player"] == selected_player].sort_values("#")
            st.dataframe(player_df.drop(columns=["Player"]), use_container_width=True, hide_index=True)

            user_id_by_username = {p["username"]: p["user_id"] for p in picks_rows}

            if st.button(f"🗑️ Remove {selected_player}'s picks for Week {view_week}", key="remove_picks_btn"):
                st.session_state.pending_removal = {
                    "week": view_week,
                    "player": selected_player,
                    "user_id": user_id_by_username.get(selected_player),
                }
                st.rerun()

    # Checked every rerun so the dialog's own buttons get a chance to be detected as clicked.
    if st.session_state.get("pending_removal"):
        confirm_removal(st.session_state.pending_removal)

    st.write("---")

    # 7. EXPORT EVERYONE'S PICKS (BEFORE GRADING)
    st.subheader("📤 Export Everyone's Picks")
    st.caption("The list you'd send out Saturday morning, before any games are graded -- just names and pick numbers.")

    picks_export_week = section_week_selector("picks_export_week", active_week)

    picks_export_games = supabase.table("games").select("*").eq("week_number", picks_export_week).execute().data
    picks_export_picks = supabase.table("picks").select("*").eq("week_number", picks_export_week).execute().data
    picks_export_players = supabase.table("players").select("*").order("username").execute().data

    if not picks_export_games:
        st.info(f"No games loaded for Week {picks_export_week} yet.")
    elif not picks_export_picks:
        st.info(f"No picks submitted yet for Week {picks_export_week}.")
    else:
        def build_picks_only_workbook(week_number, players_rows, picks_rows, games_rows):
            numbers_map = compute_game_numbers(games_rows)
            games_map = {g["game_id"]: g for g in games_rows}

            picks_by_user = {}
            for pk in picks_rows:
                g = games_map.get(pk["game_id"])
                if not g:
                    continue
                nums = numbers_map.get(pk["game_id"], {})
                is_fav = pk["selected_team"] == g.get("favorite_team")
                num = nums.get("fav_num") if is_fav else nums.get("und_num")
                if num is None:
                    continue
                picks_by_user.setdefault(pk["user_id"], []).append(num)

            wb = Workbook()
            ws = wb.active
            ws.title = f"Week {week_number} Picks"[:31]

            headers = ["Player", f"Week {week_number} Picks"]
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            rows_with_picks = [p for p in players_rows if p["id"] in picks_by_user]
            rows_with_picks.sort(key=lambda p: p["username"])

            for p in rows_with_picks:
                nums = sorted(picks_by_user[p["id"]])
                ws.append([p["username"], ",".join(str(n) for n in nums)])
                r_idx = ws.max_row
                name_cell = ws.cell(row=r_idx, column=1)
                name_cell.font = Font(name="Arial", bold=True)
                paid_color = "FF00B050" if p.get("paid") else "FFFF0000"
                name_cell.fill = PatternFill(start_color=paid_color, end_color=paid_color, fill_type="solid")
                ws.cell(row=r_idx, column=2).font = Font(name="Arial")

            for col_idx, header in enumerate(headers, start=1):
                col_letter = get_column_letter(col_idx)
                longest = max(
                    [len(str(header))] +
                    [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
                )
                ws.column_dimensions[col_letter].width = min(longest + 3, 40)

            ws.freeze_panes = "A2"

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer

        picks_only_buffer = build_picks_only_workbook(int(picks_export_week), picks_export_players, picks_export_picks, picks_export_games)
        st.download_button(
            f"⬇️ Download Week {int(picks_export_week)} Picks (.xlsx)",
            data=picks_only_buffer,
            file_name=f"week_{int(picks_export_week)}_picks.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    st.write("---")

    # 12. MANUALLY ENTER PICKS FOR A PLAYER (BY NUMBER)
    st.subheader("✍️ Manually Enter Picks for a Player")
    st.caption("For anyone who sent you picks by text instead of using the app -- enter the numbers they gave you.")

    manual_pick_week = section_week_selector("manual_pick_week", active_week)

    players_roster = supabase.table("players").select("*").execute().data
    if not players_roster:
        st.info("No players found yet -- add one above, or have them log in once.")
    else:
        player_options = {p["username"]: p["id"] for p in players_roster}
        selected_manual_player = st.selectbox(
            "Player:", ["— Select a player —"] + sorted(player_options.keys()), key="manual_pick_player"
        )
        if selected_manual_player == "— Select a player —":
            st.info("Select a player above to enter picks for them.")
        else:
            selected_manual_user_id = player_options[selected_manual_player]

            week_games_for_manual = supabase.table("games").select("*").eq("week_number", manual_pick_week).execute().data

            if not week_games_for_manual:
                st.info(f"No games loaded for Week {manual_pick_week} yet.")
            else:
                numbers_map_manual = compute_game_numbers(week_games_for_manual)
                games_by_id_manual = {g["game_id"]: g for g in week_games_for_manual}

                # number -> which game/team that number refers to, same odd/even scheme as the exports
                number_lookup = {}
                for g in week_games_for_manual:
                    nums = numbers_map_manual.get(g["game_id"], {})
                    number_lookup[nums["fav_num"]] = {"game_id": g["game_id"], "team": g.get("favorite_team", "")}
                    number_lookup[nums["und_num"]] = {"game_id": g["game_id"], "team": g.get("underdog_team", "")}

                with st.expander(f"See all game numbers for Week {manual_pick_week}"):
                    ref_rows = [
                        {
                            "#": num,
                            "Team": info["team"],
                            "Matchup": games_by_id_manual[info["game_id"]].get("display_text", ""),
                            "Spread": games_by_id_manual[info["game_id"]].get("spread_value", ""),
                        }
                        for num, info in sorted(number_lookup.items())
                    ]
                    st.dataframe(ref_rows, use_container_width=True, hide_index=True)

                existing_manual_picks = supabase.table("picks").select("*") \
                    .eq("user_id", selected_manual_user_id).eq("week_number", manual_pick_week).execute().data

                existing_numbers = []
                for pk in existing_manual_picks:
                    g = games_by_id_manual.get(pk["game_id"])
                    if not g:
                        continue
                    nums = numbers_map_manual.get(pk["game_id"], {})
                    is_fav = pk["selected_team"] == g.get("favorite_team")
                    num = nums.get("fav_num") if is_fav else nums.get("und_num")
                    if num is not None:
                        existing_numbers.append(num)
                existing_numbers.sort()

                if existing_manual_picks:
                    st.caption(f"{selected_manual_player} already has {len(existing_manual_picks)} pick(s) on file for Week {manual_pick_week} -- pre-filled below.")

                numbers_input = st.text_input(
                    "Enter their 7 picks by number, comma-separated (e.g. 3,5,11,13,41,45,144):",
                    value=",".join(str(n) for n in existing_numbers),
                    key="manual_pick_numbers_input",
                )

                if st.button(f"Save Picks for {selected_manual_player}", type="primary", key="save_manual_picks"):
                    raw_parts = [p.strip() for p in numbers_input.split(",") if p.strip()]
                    try:
                        entered_numbers = [int(p) for p in raw_parts]
                    except ValueError:
                        entered_numbers = None
                        st.error("Please enter numbers only, separated by commas.")

                    if entered_numbers is not None:
                        invalid_numbers = sorted({n for n in entered_numbers if n not in number_lookup})
                        duplicate_numbers = sorted({n for n in entered_numbers if entered_numbers.count(n) > 1})

                        if invalid_numbers:
                            st.error(f"These numbers don't match any game this week: {invalid_numbers}")
                        elif duplicate_numbers:
                            st.error(f"These numbers were entered more than once: {duplicate_numbers}")
                        elif len(entered_numbers) != 7:
                            st.error(f"You entered {len(entered_numbers)} number(s) -- exactly 7 are required.")
                        else:
                            preview_rows = [
                                {
                                    "#": n, "Team": number_lookup[n]["team"],
                                    "Matchup": games_by_id_manual[number_lookup[n]["game_id"]].get("display_text", ""),
                                    "Spread": games_by_id_manual[number_lookup[n]["game_id"]].get("spread_value", ""),
                                }
                                for n in sorted(entered_numbers)
                            ]
                            resolved = [
                                {
                                    "game_id": number_lookup[n]["game_id"],
                                    "team": number_lookup[n]["team"],
                                    "spread": games_by_id_manual[number_lookup[n]["game_id"]].get("spread_value", ""),
                                }
                                for n in entered_numbers
                            ]
                            st.session_state.pending_manual_picks_save = {
                                "week": manual_pick_week,
                                "player": selected_manual_player,
                                "user_id": selected_manual_user_id,
                                "preview_rows": preview_rows,
                                "resolved": resolved,
                            }
                            st.rerun()

    if st.session_state.get("pending_manual_picks_save"):
        confirm_manual_picks_save(st.session_state.pending_manual_picks_save)

    st.write("---")


with tab_grading:
    # 8. GRADE FINISHED GAMES
    st.subheader("🏁 Grade Finished Games")
    grade_week_num = section_week_selector("grade_week", active_week, "Week to grade:")

    if st.button("🔄 Refresh Scores & Grade", type="primary"):
        games_in_week = supabase.table("games").select("*").eq("week_number", grade_week_num).execute().data

        if not games_in_week:
            st.warning(f"No games found for Week {grade_week_num}.")
        else:
            with st.spinner("Pulling final scores from ESPN..."):
                headers = {
                    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
                    "Referer": "https://www.espn.com/",
                    "Accept": "application/json",
                }
                league_url_map = {
                    "NFL": "https://site.web.api.espn.com/apis/site/v2/sports/football/nfl/scoreboard",
                    "CFB": "https://site.web.api.espn.com/apis/site/v2/sports/football/college-football/scoreboard",
                }

                games_by_league = {}
                for g in games_in_week:
                    games_by_league.setdefault(g["league"], []).append(g)

                espn_scores = {}  # game_id -> {"state": ..., "home_score": ..., "away_score": ...}

                for league_name, league_games in games_by_league.items():
                    url = league_url_map.get(league_name)
                    if not url:
                        continue

                    kickoff_dates = []
                    for g in league_games:
                        try:
                            kickoff_dates.append(datetime.fromisoformat(g["kickoff_time"].replace("Z", "+00:00")))
                        except (ValueError, AttributeError, TypeError):
                            pass
                    if not kickoff_dates:
                        continue

                    params = {
                        "limit": 1000,
                        "dates": f"{min(kickoff_dates):%Y%m%d}-{max(kickoff_dates):%Y%m%d}",
                    }
                    if league_name == "CFB":
                        params["groups"] = 80

                    try:
                        resp = requests.get(url, params=params, headers=headers, timeout=15)
                        if resp.status_code != 200:
                            st.warning(f"{league_name} score refresh failed: HTTP {resp.status_code}")
                            continue
                        data = resp.json()
                        time_module.sleep(2)  # same rate-limit courtesy as the sync
                    except Exception as e:
                        st.warning(f"{league_name} score refresh error: {e}")
                        continue

                    for event in data.get("events", []):
                        gid = f"espn_{event.get('id')}"
                        state = event.get("status", {}).get("type", {}).get("state", "pre")

                        competitions = event.get("competitions", [{}])[0]
                        competitors = competitions.get("competitors", [])
                        home_node = next((c for c in competitors if c.get("homeAway") == "home"), {})
                        away_node = next((c for c in competitors if c.get("homeAway") == "away"), {})
                        try:
                            home_score = int(home_node.get("score", 0))
                            away_score = int(away_node.get("score", 0))
                        except (TypeError, ValueError):
                            home_score, away_score = 0, 0

                        espn_scores[gid] = {"state": state, "home_score": home_score, "away_score": away_score}

                graded_count = 0
                already_final_count = 0
                still_pending_count = 0

                for g in games_in_week:
                    if g.get("status") == "final":
                        already_final_count += 1
                        continue

                    live = espn_scores.get(g["game_id"])
                    if not live or live["state"] != "post":
                        still_pending_count += 1
                        continue

                    home_score, away_score = live["home_score"], live["away_score"]

                    # The stored spread is always relative to the FAVORITE (see nudge_off_whole_number),
                    # and every spread has been shifted off whole numbers, so a tie is not possible.
                    spread_num_str = (g.get("spread_value") or "").rsplit(" ", 1)[-1]
                    try:
                        spread_num = float(spread_num_str)
                    except ValueError:
                        spread_num = 0.0

                    fav_is_home = bool(g.get("favorite_team_home"))
                    fav_score = home_score if fav_is_home else away_score
                    und_score = away_score if fav_is_home else home_score

                    margin = (fav_score + spread_num) - und_score
                    winning_team = g.get("favorite_team") if margin > 0 else g.get("underdog_team")

                    supabase.table("games").update({
                        "status": "final",
                        "winning_team": winning_team,
                    }).eq("id", g["id"]).execute()

                    # Grade each pick against the spread IT actually saw, not necessarily
                    # today's spread_value -- protects anyone who picked before a re-sync.
                    picks_for_game = supabase.table("picks").select("*").eq("game_id", g["game_id"]).execute().data
                    for p in picks_for_game:
                        locked_spread_str = (p.get("spread_at_pick") or g.get("spread_value") or "").rsplit(" ", 1)[-1]
                        try:
                            locked_spread_num = float(locked_spread_str)
                        except ValueError:
                            locked_spread_num = spread_num

                        pick_margin = (fav_score + locked_spread_num) - und_score
                        pick_winning_team = g.get("favorite_team") if pick_margin > 0 else g.get("underdog_team")
                        pick_result = "win" if p.get("selected_team") == pick_winning_team else "loss"

                        supabase.table("picks").update({"result": pick_result}).eq("id", p["id"]).execute()

                    graded_count += 1

            st.success(
                f"Graded {graded_count} newly-final game(s) for Week {grade_week_num}. "
                f"{already_final_count} were already graded, {still_pending_count} still in progress or not yet started."
            )

    st.write("---")


with tab_players:
    # 11. ADD A PLAYER WITHOUT AN ACCOUNT
    st.subheader("➕ Add a Player Without an Account")
    st.caption("For someone you're tracking who doesn't log in or have an email on file.")

    manual_player_name = st.text_input("Player name:", key="manual_player_name")
    if st.button("Add Player", key="add_manual_player_btn"):
        if not manual_player_name.strip():
            st.error("Enter a name.")
        else:
            try:
                supabase.table("players").insert({
                    "id": str(uuid.uuid4()),
                    "username": manual_player_name.strip(),
                }).execute()
                st.success(f"Added {manual_player_name.strip()}.")
                st.rerun()
            except Exception as e:
                st.error(f"Database error: {e}")

    st.write("---")

    # 11a. BULK ADD PLAYERS WITHOUT ACCOUNTS
    st.subheader("📋 Bulk Add Players")
    st.caption("Paste one name per line to add several at once -- e.g. importing a whole roster from a spreadsheet.")

    bulk_names_text = st.text_area(
        "One player name per line:",
        height=150,
        key="bulk_player_names",
        placeholder="Trent Y.\nDusty Y.\nJoe P.\nJake H.",
    )

    if st.button("Add All", key="bulk_add_players_btn"):
        raw_names = [n.strip() for n in bulk_names_text.split("\n") if n.strip()]
        if not raw_names:
            st.error("Paste at least one name.")
        else:
            existing_players = supabase.table("players").select("username").execute().data
            existing_usernames = {p["username"] for p in existing_players}

            added, skipped = [], []
            for name in raw_names:
                if name in existing_usernames:
                    skipped.append(name)
                    continue
                try:
                    supabase.table("players").insert({
                        "id": str(uuid.uuid4()),
                        "username": name,
                    }).execute()
                    added.append(name)
                    existing_usernames.add(name)  # guard against duplicate lines in the same paste
                except Exception as e:
                    st.error(f"Failed to add {name}: {e}")

            if added:
                st.success(f"Added {len(added)} player(s): {', '.join(added)}")
            if skipped:
                st.info(f"Skipped {len(skipped)} already-existing name(s): {', '.join(skipped)}")
            st.rerun()

    st.write("---")

    # 13. MERGE A MANUAL PLAYER INTO A REAL ACCOUNT
    st.subheader("🔀 Merge a Manual Player into a Real Account")
    st.caption("Once someone you added manually creates a real login, combine their pick history under the real account.")

    merge_players_roster = supabase.table("players").select("*").order("username").execute().data
    if len(merge_players_roster) < 2:
        st.info("Need at least two player entries to merge.")
    else:
        merge_options = {p["username"]: p["id"] for p in merge_players_roster}
        col_from, col_to = st.columns(2)
        with col_from:
            merge_from_name = st.selectbox(
                "Merge FROM (the manual entry):", ["— Select a player —"] + sorted(merge_options.keys()), key="merge_from"
            )
        with col_to:
            to_choices = ["— Select a player —"] + [n for n in sorted(merge_options.keys()) if n != merge_from_name]
            merge_to_name = st.selectbox("Merge INTO (the real account):", to_choices, key="merge_to")

        if merge_from_name == "— Select a player —" or merge_to_name == "— Select a player —":
            st.info("Select both a FROM and an INTO player above to merge.")
        elif st.button("Merge Players", type="primary", key="merge_players_btn"):
            from_id = merge_options[merge_from_name]
            to_id = merge_options[merge_to_name]

            from_picks = supabase.table("picks").select("*").eq("user_id", from_id).execute().data
            to_picks = supabase.table("picks").select("week_number, game_id").eq("user_id", to_id).execute().data

            to_game_ids_by_week = {}
            for p in to_picks:
                to_game_ids_by_week.setdefault(p["week_number"], set()).add(p["game_id"])

            conflict_weeks = sorted({
                p["week_number"] for p in from_picks
                if p["game_id"] in to_game_ids_by_week.get(p["week_number"], set())
            })

            st.session_state.pending_merge = {
                "from_id": from_id, "from_name": merge_from_name,
                "to_id": to_id, "to_name": merge_to_name,
                "from_picks_count": len(from_picks),
                "conflict_weeks": conflict_weeks,
                "completed": False,
            }
            st.rerun()

    if st.session_state.get("pending_merge"):
        confirm_merge(st.session_state.pending_merge)

    st.write("---")

    # 14. PLAYER PAYMENT STATUS
    st.subheader("💰 Player Payment Status")

    players_for_payment = supabase.table("players").select("*").order("username").execute().data
    if not players_for_payment:
        st.info("No players yet -- they need to log in at least once before they show up here.")
    else:
        payment_df = pd.DataFrame([
            {"Player": p["username"], "Paid": bool(p.get("paid")), "_id": p["id"]}
            for p in players_for_payment
        ])
        edited_payment_df = st.data_editor(
            payment_df.drop(columns=["_id"]),
            column_config={"Paid": st.column_config.CheckboxColumn("Paid?")},
            hide_index=True,
            use_container_width=True,
            key="payment_editor",
        )
        if st.button("Save Payment Status", key="save_payment_status"):
            try:
                for i, row in edited_payment_df.iterrows():
                    player_id = payment_df.iloc[i]["_id"]
                    supabase.table("players").update({"paid": bool(row["Paid"])}).eq("id", player_id).execute()
                st.success("Payment status updated.")
            except Exception as e:
                st.error(f"Database error: {e}")

    st.write("---")

    # 15. DELETE A PLAYER
    st.subheader("🗑️ Delete a Player")
    st.caption("For a player added by mistake, a duplicate, or someone who's leaving the league entirely.")

    players_for_delete = supabase.table("players").select("*").order("username").execute().data
    if not players_for_delete:
        st.info("No players found yet.")
    else:
        delete_player_options = {p["username"]: p["id"] for p in players_for_delete}
        selected_delete_player = st.selectbox(
            "Player to delete:", ["— Select a player —"] + sorted(delete_player_options.keys()), key="delete_player_select"
        )

        if selected_delete_player == "— Select a player —":
            st.info("Select a player above to delete them.")
        else:
            selected_delete_id = delete_player_options[selected_delete_player]

            if st.button(f"Delete {selected_delete_player}", key="delete_player_btn"):
                picks_for_player = supabase.table("picks").select("week_number").eq("user_id", selected_delete_id).execute().data
                st.session_state.pending_player_delete = {
                    "player_id": selected_delete_id,
                    "name": selected_delete_player,
                    "picks_count": len(picks_for_player),
                    "weeks_count": len({p["week_number"] for p in picks_for_player}),
                }
                st.rerun()

    if st.session_state.get("pending_player_delete"):
        confirm_delete_player(st.session_state.pending_player_delete)

    st.write("---")


with tab_exports:
    # 9. EXPORT SLATE TO EXCEL
    st.subheader("📊 Export Slate to Excel")
    export_week = section_week_selector("export_week", active_week, "Week to export:")

    export_games = supabase.table("games").select("*").eq("week_number", export_week).execute().data

    if not export_games:
        st.info(f"No games found for Week {export_week}.")
    else:
        def build_slate_workbook(games_rows, week_number):
            wb = Workbook()
            ws = wb.active
            ws.title = f"Week {week_number}"[:31]

            headers = ["#", "FAVORITE", "#", "UNDERDOG", "SPREAD", "KICKOFF (ET)", "TV"]

            ws.append([f"Week {week_number}"])
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
            title_cell = ws.cell(row=1, column=1)
            title_cell.font = Font(name="Arial", bold=True, size=14)
            title_cell.alignment = Alignment(horizontal="center")

            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=2, column=col_idx)
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            sorted_games = sorted(games_rows, key=lambda g: g.get("kickoff_time") or "")
            for i, g in enumerate(sorted_games, start=1):
                fav_num = 2 * i - 1  # favorites get odd numbers
                und_num = 2 * i      # underdogs get even numbers

                fav_team = g.get("favorite_team", "")
                und_team = g.get("underdog_team", "")
                fav_team = f"{fav_team} (Home)" if g.get("favorite_team_home") else fav_team
                und_team = f"{und_team} (Home)" if g.get("underdog_team_home") else und_team
                spread_number = (g.get("spread_value") or "").rsplit(" ", 1)[-1]

                kickoff_display = g.get("kickoff_time", "") or ""
                try:
                    kickoff_dt = datetime.fromisoformat(kickoff_display.replace("Z", "+00:00")).astimezone(ZoneInfo("America/New_York"))
                    kickoff_display = kickoff_dt.strftime("%a %m/%d %I:%M %p ET").replace(" 0", " ")
                except (ValueError, AttributeError):
                    pass

                tv_network = g.get("tv_network", "") or ""

                ws.append([fav_num, fav_team, und_num, und_team, spread_number, kickoff_display, tv_network])
                for col_idx in range(1, len(headers) + 1):
                    ws.cell(row=i + 2, column=col_idx).font = Font(name="Arial")

            for col_idx, header in enumerate(headers, start=1):
                col_letter = get_column_letter(col_idx)
                longest = max(
                    [len(str(header))] +
                    [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(3, ws.max_row + 1)]
                )
                ws.column_dimensions[col_letter].width = longest + 4

            ws.freeze_panes = "A3"

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer

        excel_buffer = build_slate_workbook(export_games, export_week)
        st.download_button(
            "⬇️ Download Week's Slate (.xlsx)",
            data=excel_buffer,
            file_name=f"week_{export_week}_slate.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    st.write("---")

    # 16. EXPORT SEASON TRACKER
    st.subheader("📥 Export Season Tracker (.xlsx)")
    st.caption("A running week-by-week win/loss breakdown for every player, in the same style as your old sheet.")

    all_week_status_rows = supabase.table("games").select("week_number, status").execute().data
    graded_weeks = sorted({g["week_number"] for g in all_week_status_rows if g.get("status") == "final"})

    if not graded_weeks:
        st.info("No graded weeks yet -- use 'Grade Finished Games' above first.")
    else:
        players_for_export = supabase.table("players").select("*").order("username").execute().data
        if not players_for_export:
            st.info("No players found yet.")
        else:
            def build_season_tracker_workbook(players_rows, weeks):
                wb = Workbook()
                ws = wb.active
                ws.title = "Season Tracker"

                headers = ["Player"]
                for wk in weeks:
                    headers += [f"Week {wk} Picks", f"Week {wk} Wins", f"Week {wk} Loss", f"Week {wk} W", f"Week {wk} L"]
                headers += ["Overall Wins", "Overall Losses", "Place"]

                ws.append(headers)
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center")

                # Pre-fetch each week's games/picks once
                week_game_numbers, week_games_by_id, week_picks = {}, {}, {}
                for wk in weeks:
                    games_this_week = supabase.table("games").select("*").eq("week_number", wk).execute().data
                    week_game_numbers[wk] = compute_game_numbers(games_this_week)
                    week_games_by_id[wk] = {g["game_id"]: g for g in games_this_week}
                    week_picks[wk] = supabase.table("picks").select("*").eq("week_number", wk).execute().data

                # Build every player's row data first so ranking can happen before writing
                player_rows_data = []
                for p in players_rows:
                    weekly_cells = []
                    weekly_raw = []  # one entry per week: (picks_nums, win_nums, loss_nums), for coloring after append
                    total_wins, total_losses = 0, 0

                    for wk in weeks:
                        games_map = week_games_by_id[wk]
                        numbers_map = week_game_numbers[wk]
                        player_picks = [pk for pk in week_picks[wk] if pk["user_id"] == p["id"]]

                        picks_nums, win_nums, loss_nums = [], [], []
                        for pk in player_picks:
                            g = games_map.get(pk["game_id"])
                            if not g:
                                continue
                            nums = numbers_map.get(pk["game_id"], {})
                            is_fav = pk["selected_team"] == g.get("favorite_team")
                            num = nums.get("fav_num") if is_fav else nums.get("und_num")
                            if num is None:
                                continue
                            picks_nums.append(num)
                            if pk.get("result") == "win":
                                win_nums.append(num)
                            elif pk.get("result") == "loss":
                                loss_nums.append(num)

                        picks_nums.sort(); win_nums.sort(); loss_nums.sort()
                        weekly_raw.append((picks_nums, win_nums, loss_nums))
                        weekly_cells += [
                            None,  # Picks -- filled in with colored rich text after the row is appended
                            ",".join(str(n) for n in win_nums),
                            ",".join(str(n) for n in loss_nums),
                            len(win_nums),
                            len(loss_nums),
                        ]
                        total_wins += len(win_nums)
                        total_losses += len(loss_nums)

                    player_rows_data.append({
                        "username": p["username"],
                        "paid": bool(p.get("paid")),
                        "weekly_cells": weekly_cells,
                        "weekly_raw": weekly_raw,
                        "total_wins": total_wins,
                        "total_losses": total_losses,
                    })

                # Competition ranking -- ties share a place, same as your RANK() formula
                for row in player_rows_data:
                    row["place"] = sum(1 for r in player_rows_data if r["total_wins"] > row["total_wins"]) + 1
                player_rows_data.sort(key=lambda r: (-r["total_wins"], r["username"]))

                for row in player_rows_data:
                    ws.append([row["username"]] + row["weekly_cells"] + [row["total_wins"], row["total_losses"], row["place"]])
                    r_idx = ws.max_row
                    name_cell = ws.cell(row=r_idx, column=1)
                    name_cell.font = Font(name="Arial", bold=True)
                    paid_color = "FF00B050" if row["paid"] else "FFFF0000"
                    name_cell.fill = PatternFill(start_color=paid_color, end_color=paid_color, fill_type="solid")

                    for week_idx, (picks_nums, win_nums, loss_nums) in enumerate(row["weekly_raw"]):
                        picks_col = 2 + week_idx * 5  # 1-based: Player is col 1, then 5 cols per week
                        ws.cell(row=r_idx, column=picks_col).value = colored_picks_richtext(picks_nums, win_nums, loss_nums)
                        for offset in (1, 2, 3, 4):  # Wins, Loss, W, L -- everything except the Picks column
                            ws.cell(row=r_idx, column=picks_col + offset).font = Font(name="Arial")

                    # Overall Wins / Overall Losses / Place columns, at the very end
                    for col_idx in range(2 + len(row["weekly_raw"]) * 5, len(headers) + 1):
                        ws.cell(row=r_idx, column=col_idx).font = Font(name="Arial")

                for col_idx in range(1, len(headers) + 1):
                    col_letter = get_column_letter(col_idx)
                    longest = max(
                        [len(str(headers[col_idx - 1]))] +
                        [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
                    )
                    ws.column_dimensions[col_letter].width = min(longest + 3, 22)

                ws.freeze_panes = "B2"

                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                return buffer

            tracker_buffer = build_season_tracker_workbook(players_for_export, graded_weeks)
            st.download_button(
                "⬇️ Download Season Tracker (.xlsx)",
                data=tracker_buffer,
                file_name="season_tracker.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    st.write("---")

    # 17. EXPORT WEEK + SEASON RESULTS (what you'd send out to the group)
    st.subheader("📤 Export Week + Season Results")
    st.caption("What you'd send out after grading a week: that week's individual results, plus updated season standings.")

    results_week = section_week_selector("results_export_week", active_week, "Week to report on:")

    if results_week not in graded_weeks:
        st.info(f"Week {int(results_week)} hasn't been graded yet -- use 'Grade Finished Games' above first.")
    else:
        players_for_results = supabase.table("players").select("*").order("username").execute().data
        if not players_for_results:
            st.info("No players found yet.")
        else:
            def build_week_and_season_workbook(players_rows, target_week, all_graded_weeks):
                wb = Workbook()

                def style_header(ws, headers):
                    ws.append(headers)
                    for col_idx in range(1, len(headers) + 1):
                        cell = ws.cell(row=1, column=col_idx)
                        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center")

                def autosize(ws, headers):
                    for col_idx, header in enumerate(headers, start=1):
                        col_letter = get_column_letter(col_idx)
                        longest = max(
                            [len(str(header))] +
                            [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
                        )
                        ws.column_dimensions[col_letter].width = min(longest + 3, 22)
                    ws.freeze_panes = "A2"

                # --- Sheet 1: this week's individual results ---
                ws1 = wb.active
                ws1.title = f"Week {target_week} Results"[:31]

                week_games = supabase.table("games").select("*").eq("week_number", target_week).execute().data
                numbers_map = compute_game_numbers(week_games)
                games_map = {g["game_id"]: g for g in week_games}
                week_picks_rows = supabase.table("picks").select("*").eq("week_number", target_week).execute().data

                headers1 = ["Player", "Picks", "Wins", "Loss", "W", "L"]
                style_header(ws1, headers1)

                week_rows_data = []
                for p in players_rows:
                    player_picks = [pk for pk in week_picks_rows if pk["user_id"] == p["id"]]
                    if not player_picks:
                        continue
                    picks_nums, win_nums, loss_nums = [], [], []
                    for pk in player_picks:
                        g = games_map.get(pk["game_id"])
                        if not g:
                            continue
                        nums = numbers_map.get(pk["game_id"], {})
                        is_fav = pk["selected_team"] == g.get("favorite_team")
                        num = nums.get("fav_num") if is_fav else nums.get("und_num")
                        if num is None:
                            continue
                        picks_nums.append(num)
                        if pk.get("result") == "win":
                            win_nums.append(num)
                        elif pk.get("result") == "loss":
                            loss_nums.append(num)
                    picks_nums.sort(); win_nums.sort(); loss_nums.sort()
                    week_rows_data.append({
                        "username": p["username"], "picks": picks_nums, "wins": win_nums,
                        "losses": loss_nums, "w": len(win_nums), "l": len(loss_nums),
                    })

                week_rows_data.sort(key=lambda r: (-r["w"], r["l"], r["username"]))
                for r in week_rows_data:
                    ws1.append([
                        r["username"], None, ",".join(map(str, r["wins"])),
                        ",".join(map(str, r["losses"])), r["w"], r["l"],
                    ])
                    row_idx = ws1.max_row
                    ws1.cell(row=row_idx, column=2).value = colored_picks_richtext(r["picks"], r["wins"], r["losses"])
                    for col_idx in (1, 3, 4, 5, 6):
                        ws1.cell(row=row_idx, column=col_idx).font = Font(name="Arial")

                autosize(ws1, headers1)

                # --- Sheet 2: season standings ---
                ws2 = wb.create_sheet(title="Season Standings")

                header_font = Font(name="Arial", bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                center = Alignment(horizontal="center")
                divider = Border(right=Side(style="thin", color="000000"))

                ws2.cell(row=1, column=1, value="Player")
                ws2.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
                ws2.cell(row=1, column=2, value="Overall Record")
                ws2.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
                ws2.cell(row=1, column=4, value="Place")
                ws2.merge_cells(start_row=1, start_column=4, end_row=2, end_column=4)
                ws2.cell(row=1, column=5, value="Paid")
                ws2.merge_cells(start_row=1, start_column=5, end_row=2, end_column=5)

                for hdr_row in (1, 2):
                    for col_idx in range(1, 6):
                        cell = ws2.cell(row=hdr_row, column=col_idx)
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = center

                season_totals = {}
                for wk in all_graded_weeks:
                    picks_wk = supabase.table("picks").select("user_id, result").eq("week_number", wk).execute().data
                    for pk in picks_wk:
                        uid = pk["user_id"]
                        totals = season_totals.setdefault(uid, {"wins": 0, "losses": 0})
                        if pk.get("result") == "win":
                            totals["wins"] += 1
                        elif pk.get("result") == "loss":
                            totals["losses"] += 1

                season_rows = []
                for p in players_rows:
                    totals = season_totals.get(p["id"], {"wins": 0, "losses": 0})
                    wins, losses = totals["wins"], totals["losses"]
                    season_rows.append({
                        "username": p["username"], "wins": wins, "losses": losses,
                        "paid": bool(p.get("paid")),
                    })

                for row in season_rows:
                    row["place"] = sum(1 for r in season_rows if r["wins"] > row["wins"]) + 1
                season_rows.sort(key=lambda r: (-r["wins"], r["losses"], r["username"]))

                medal_colors = {1: "FFD700", 2: "C0C0C0", 3: "CD7F32"}

                for row in season_rows:
                    ws2.append([row["username"], row["wins"], row["losses"], row["place"], "Yes" if row["paid"] else "No"])
                    r_idx = ws2.max_row
                    name_cell = ws2.cell(row=r_idx, column=1)
                    name_cell.font = Font(name="Arial", bold=True)
                    paid_color = "FF00B050" if row["paid"] else "FFFF0000"
                    name_cell.fill = PatternFill(start_color=paid_color, end_color=paid_color, fill_type="solid")
                    for col_idx in (2, 3, 4, 5):
                        ws2.cell(row=r_idx, column=col_idx).font = Font(name="Arial")
                    ws2.cell(row=r_idx, column=2).border = divider  # vertical line between wins/losses

                    if row["place"] in medal_colors:
                        place_cell = ws2.cell(row=r_idx, column=4)
                        place_cell.font = Font(name="Arial", bold=True)
                        place_color = medal_colors[row["place"]]
                        place_cell.fill = PatternFill(start_color=place_color, end_color=place_color, fill_type="solid")

                headers2_display = ["Player", "Overall Record", "", "Place", "Paid"]
                for col_idx, header in enumerate(headers2_display, start=1):
                    col_letter = get_column_letter(col_idx)
                    longest = max(
                        [len(str(header))] +
                        [len(str(ws2.cell(row=r, column=col_idx).value or "")) for r in range(3, ws2.max_row + 1)]
                    )
                    ws2.column_dimensions[col_letter].width = min(longest + 3, 22)
                ws2.freeze_panes = "A3"

                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                return buffer

            results_buffer = build_week_and_season_workbook(players_for_results, int(results_week), graded_weeks)
            st.download_button(
                f"⬇️ Download Week {int(results_week)} + Season Results (.xlsx)",
                data=results_buffer,
                file_name=f"week_{int(results_week)}_and_season_results.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

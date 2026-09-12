import streamlit as st
import pandas as pd
import io
import requests
import time as time_module
from datetime import datetime
from streamlit_autorefresh import st_autorefresh
from supabase import create_client, Client
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SUPABASE_URL = st.secrets.get("SUPABASE_URL")
SUPABASE_KEY = st.secrets.get("SUPABASE_KEY")
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

if st.session_state.get("access_token"):
    try:
        supabase.auth.set_session(st.session_state.access_token, st.session_state.refresh_token)
    except Exception:
        pass

is_admin = False
if st.session_state.get("user"):
    try:
        role_resp = supabase.table("league_users").select("role").eq("id", st.session_state.user.id).execute()
        if role_resp.data and role_resp.data[0].get("role") == "admin":
            is_admin = True
    except Exception:
        pass

st.title("🏆 League Standings & Leaderboard")
st.write("Running tab of each participant's record week over week.")

# Refresh controls, same idea as the main app page -- available to everyone,
# since seeing updated standings matters to players too, not just admins.
col_refresh_btn, col_auto_toggle, col_auto_interval = st.columns([1, 1, 1])
with col_refresh_btn:
    if st.button("🔄 Refresh Now"):
        st.rerun()
with col_auto_toggle:
    lb_auto_refresh_on = st.checkbox("Auto-refresh", key="lb_auto_refresh_enabled")
with col_auto_interval:
    if lb_auto_refresh_on:
        lb_interval_label = st.selectbox(
            "Every:", ["60 sec", "2 min", "5 min"], key="lb_auto_refresh_interval", label_visibility="collapsed",
        )
        lb_interval_seconds = {"60 sec": 60, "2 min": 120, "5 min": 300}[lb_interval_label]
        st_autorefresh(interval=lb_interval_seconds * 1000, key="lb_autorefresh")

if is_admin:
    with st.expander("🏁 Grade Finished Games (Admin)"):
        # Follows the same shared "global_week" used on the app and Admin pages --
        # changing it here also updates it there, and vice versa.
        if "global_week" not in st.session_state:
            st.session_state["global_week"] = 1

        if st.session_state.get("_lb_last_seen_global") != st.session_state["global_week"]:
            st.session_state["lb_grade_week"] = st.session_state["global_week"]
            st.session_state["_lb_last_seen_global"] = st.session_state["global_week"]

        if "lb_grade_week" not in st.session_state:
            st.session_state["lb_grade_week"] = st.session_state["global_week"]

        grade_week_num = st.number_input("Week to grade:", min_value=1, max_value=18, step=1, key="lb_grade_week")

        if grade_week_num != st.session_state["global_week"]:
            st.session_state["global_week"] = grade_week_num
            st.session_state["_lb_last_seen_global"] = grade_week_num

        if st.button("🔄 Refresh Scores & Grade", type="primary", key="lb_grade_btn"):
            games_in_week = supabase.table("games").select("*").eq("week_number", grade_week_num).execute().data

            if not games_in_week:
                st.warning(f"No games found for Week {grade_week_num}.")
            else:
                with st.spinner("Pulling final scores from ESPN..."):
                    headers = {
                        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
                        "Referer": "https://www.espn.com/",
                        "Accept": "application/json",
                    }
                    league_url_map = {
                        "NFL": "https://site.web.api.espn.com/apis/site/v2/sports/football/nfl/scoreboard",
                        "CFB": "https://site.web.api.espn.com/apis/site/v2/sports/football/college-football/scoreboard",
                    }

                    games_by_league = {}
                    for g in games_in_week:
                        games_by_league.setdefault(g["league"], []).append(g)

                    espn_scores = {}

                    for league_name, league_games in games_by_league.items():
                        url = league_url_map.get(league_name)
                        if not url:
                            continue

                        kickoff_dates = []
                        for g in league_games:
                            try:
                                kickoff_dates.append(datetime.fromisoformat(g["kickoff_time"].replace("Z", "+00:00")))
                            except (ValueError, AttributeError, TypeError):
                                pass
                        if not kickoff_dates:
                            continue

                        params = {
                            "limit": 1000,
                            "dates": f"{min(kickoff_dates):%Y%m%d}-{max(kickoff_dates):%Y%m%d}",
                        }
                        if league_name == "CFB":
                            params["groups"] = 80

                        try:
                            resp = requests.get(url, params=params, headers=headers, timeout=15)
                            if resp.status_code != 200:
                                st.warning(f"{league_name} score refresh failed: HTTP {resp.status_code}")
                                continue
                            data = resp.json()
                            time_module.sleep(2)
                        except Exception as e:
                            st.warning(f"{league_name} score refresh error: {e}")
                            continue

                        for event in data.get("events", []):
                            gid = f"espn_{event.get('id')}"
                            state = event.get("status", {}).get("type", {}).get("state", "pre")

                            competitions = event.get("competitions", [{}])[0]
                            competitors = competitions.get("competitors", [])
                            home_node = next((c for c in competitors if c.get("homeAway") == "home"), {})
                            away_node = next((c for c in competitors if c.get("homeAway") == "away"), {})
                            try:
                                home_score = int(home_node.get("score", 0))
                                away_score = int(away_node.get("score", 0))
                            except (TypeError, ValueError):
                                home_score, away_score = 0, 0

                            espn_scores[gid] = {"state": state, "home_score": home_score, "away_score": away_score}

                    graded_count = 0
                    already_final_count = 0
                    still_pending_count = 0

                    for g in games_in_week:
                        if g.get("status") == "final":
                            already_final_count += 1
                            continue

                        live = espn_scores.get(g["game_id"])
                        if not live or live["state"] != "post":
                            still_pending_count += 1
                            continue

                        home_score, away_score = live["home_score"], live["away_score"]

                        spread_num_str = (g.get("spread_value") or "").rsplit(" ", 1)[-1]
                        try:
                            spread_num = float(spread_num_str)
                        except ValueError:
                            spread_num = 0.0

                        fav_is_home = bool(g.get("favorite_team_home"))
                        fav_score = home_score if fav_is_home else away_score
                        und_score = away_score if fav_is_home else home_score

                        margin = (fav_score + spread_num) - und_score
                        winning_team = g.get("favorite_team") if margin > 0 else g.get("underdog_team")

                        supabase.table("games").update({
                            "status": "final",
                            "winning_team": winning_team,
                        }).eq("id", g["id"]).execute()

                        picks_for_game = supabase.table("picks").select("*").eq("game_id", g["game_id"]).execute().data
                        for p in picks_for_game:
                            locked_spread_str = (p.get("spread_at_pick") or g.get("spread_value") or "").rsplit(" ", 1)[-1]
                            try:
                                locked_spread_num = float(locked_spread_str)
                            except ValueError:
                                locked_spread_num = spread_num

                            pick_margin = (fav_score + locked_spread_num) - und_score
                            pick_winning_team = g.get("favorite_team") if pick_margin > 0 else g.get("underdog_team")
                            pick_result = "win" if p.get("selected_team") == pick_winning_team else "loss"

                            supabase.table("picks").update({"result": pick_result}).eq("id", p["id"]).execute()

                        graded_count += 1

                st.success(
                    f"Graded {graded_count} newly-final game(s) for Week {grade_week_num}. "
                    f"{already_final_count} were already graded, {still_pending_count} still in progress or not yet started."
                )
                st.rerun()

try:
    # Pull every pick's graded result -- this is now written directly onto each pick
    # by the Admin "Grade Finished Games" step, using the spread that pick actually
    # locked in, so it stays correct even if a week's games get re-synced later.
    picks_data = supabase.table("picks").select("username, result").execute().data

    if not picks_data:
        st.info("No player picks have been submitted yet in this league.")
    else:
        df_picks = pd.DataFrame(picks_data)
        df_graded = df_picks[df_picks["result"].isin(["win", "loss"])]

        if df_graded.empty:
            st.info("🏈 Total standings will update here as soon as the commissioner grades completed games.")
            st.subheader("Active League Roster")
            st.dataframe(df_picks[["username"]].drop_duplicates().reset_index(drop=True), use_container_width=True)
        else:
            standings = df_graded.groupby("username")["result"].agg(
                Wins=lambda r: (r == "win").sum(),
                Losses=lambda r: (r == "loss").sum(),
            ).reset_index()
            standings["Win %"] = (standings["Wins"] / (standings["Wins"] + standings["Losses"])).round(3)

            # Rank by total WINS (not win %) -- someone who skips a week shouldn't
            # rank higher just for having fewer picks. Ties share the same place,
            # same as "1-2-2-4" scoring: method="min" matches Excel's RANK().
            standings["Place"] = standings["Wins"].rank(method="min", ascending=False).astype(int)
            standings = standings.sort_values(["Wins", "Losses"], ascending=[False, True]).reset_index(drop=True)

            medals = {1: "🥇", 2: "🥈", 3: "🥉"}
            standings["Rank"] = standings["Place"].map(lambda p: medals.get(p, str(p)))
            standings = standings.rename(columns={"username": "Player"})

            display_standings = standings.copy()
            display_standings["Overall Record"] = display_standings["Wins"].astype(str) + " | " + display_standings["Losses"].astype(str)
            display_standings = display_standings[["Rank", "Player", "Overall Record", "Win %"]]

            st.subheader("🔥 Current Standings")
            st.dataframe(display_standings, use_container_width=True, hide_index=True)

            if is_admin:
                players_rows = supabase.table("players").select("username, paid").execute().data
                paid_by_username = {p["username"]: bool(p.get("paid")) for p in players_rows}

                def build_leaderboard_workbook(standings_df):
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Season Standings"

                    header_font = Font(name="Arial", bold=True, color="FFFFFF")
                    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    center = Alignment(horizontal="center")
                    divider = Border(right=Side(style="thin", color="000000"))

                    # Row 1: Player and Place/Paid span two rows; Overall Record spans two columns.
                    ws.cell(row=1, column=1, value="Player")
                    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
                    ws.cell(row=1, column=2, value="Overall Record")
                    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
                    ws.cell(row=1, column=4, value="Place")
                    ws.merge_cells(start_row=1, start_column=4, end_row=2, end_column=4)
                    ws.cell(row=1, column=5, value="Paid")
                    ws.merge_cells(start_row=1, start_column=5, end_row=2, end_column=5)

                    for row in (1, 2):
                        for col_idx in range(1, 6):
                            cell = ws.cell(row=row, column=col_idx)
                            cell.font = header_font
                            cell.fill = header_fill
                            cell.alignment = center

                    medal_colors = {1: "FFD700", 2: "C0C0C0", 3: "CD7F32"}

                    for _, row in standings_df.iterrows():
                        paid = paid_by_username.get(row["Player"], False)
                        ws.append([row["Player"], int(row["Wins"]), int(row["Losses"]), row["Rank"], "Yes" if paid else "No"])
                        r_idx = ws.max_row
                        name_cell = ws.cell(row=r_idx, column=1)
                        name_cell.font = Font(name="Arial", bold=True)
                        paid_color = "FF00B050" if paid else "FFFF0000"
                        name_cell.fill = PatternFill(start_color=paid_color, end_color=paid_color, fill_type="solid")
                        for col_idx in (2, 3, 4, 5):
                            ws.cell(row=r_idx, column=col_idx).font = Font(name="Arial")
                        ws.cell(row=r_idx, column=2).border = divider  # the vertical line between wins/losses

                        if int(row["Place"]) in medal_colors:
                            rank_cell = ws.cell(row=r_idx, column=4)
                            rank_cell.font = Font(name="Arial", bold=True)
                            rank_color = medal_colors[int(row["Place"])]
                            rank_cell.fill = PatternFill(start_color=rank_color, end_color=rank_color, fill_type="solid")

                    headers = ["Player", "Overall Record", "", "Place", "Paid"]
                    for col_idx, header in enumerate(headers, start=1):
                        col_letter = get_column_letter(col_idx)
                        longest = max(
                            [len(str(header))] +
                            [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(3, ws.max_row + 1)]
                        )
                        ws.column_dimensions[col_letter].width = min(longest + 3, 22)
                    ws.freeze_panes = "A3"

                    buffer = io.BytesIO()
                    wb.save(buffer)
                    buffer.seek(0)
                    return buffer

                leaderboard_buffer = build_leaderboard_workbook(standings)
                st.download_button(
                    "⬇️ Export Leaderboard (.xlsx)",
                    data=leaderboard_buffer,
                    file_name="season_leaderboard.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

except Exception as e:
    st.error(f"Error building standings board: {e}")
